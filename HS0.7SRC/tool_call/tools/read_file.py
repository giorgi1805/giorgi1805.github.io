import fitz 
import tiktoken
import openpyxl
from docx import Document

def context(text): 
    return text if len(tiktoken.encoding_for_model("gpt-5").encode(text)) < 100000 else f"Error: context too big, {len(text)} symbols" 

def pdf(file):  
    text = ""
    for page in fitz.open(file):
        text += page.get_text()

    return context(text)

def txt(file): 
    with open(file, "r", encoding="utf-8") as f:
        return context(f.read())

def excel(file):
    return context(openpyxl.load_workbook(file))

def word(file):
    return context("\n".join(p.text for p in Document(file).paragraphs))

def read(file): 
    typ = file.split(".")[-1]

    for xl in ['xlsx', 'xls']:
        if typ == xl: 
            return excel(file) 

    if typ == "pdf": 
        return pdf(file) 
    elif typ =="docx": 
        return word(file)

    try: 
        return txt(file)
    except UnicodeDecodeError: 
        return f"Unknown file format {typ}"

def run(file): 
    text = read(file)

    text = text.split("\n") 
    
    result = ""

    line_count = 1
    for line in text: 
        result += str(line_count) + ((5 - len(str(line_count))) * " ") + line + "\n"
        line_count += 1

    return result